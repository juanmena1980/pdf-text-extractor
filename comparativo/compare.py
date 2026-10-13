from __future__ import annotations

import difflib
import os
import re
import shutil
import unicodedata
import urllib.error
import urllib.request
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urlparse

from extract_pdf_text import extract_one_pdf


try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    raise SystemExit(
        "openpyxl no esta instalado. Ejecuta: python -m pip install -r requirements.txt"
    ) from exc


_LINK_HINTS = (
    "liga",
    "hipervinculo",
    "enlace",
    "url",
    "link",
    "pdf",
    "ruta",
    "archivo",
)
_TEXT_HINTS = (
    "transcripcion revisada",
    "texto revisado",
    "version revisada",
    "transcripcion",
    "texto",
    "contenido",
    "gold",
    "expected",
    "nota",
)
_HYPERLINK_URL = re.compile(
    r"""^=HYPERLINK\(\s*(?:"(?P<q>[^"]*)"|'(?P<s>[^']*)')""",
    re.IGNORECASE,
)
_HYPERLINK_LABEL = re.compile(
    r"""^=HYPERLINK\(\s*(?:"[^"]*"|'[^']*')\s*,\s*(?:"(?P<q>[^"]*)"|'(?P<s>[^']*)')""",
    re.IGNORECASE,
)
_ZERO_WIDTH = re.compile(r"[\u00ad\u200b\u200c\u200d\ufeff]")
_MAX_DATA_ROWS = 5000
_MAX_PDF_BYTES = 80 * 1024 * 1024


@dataclass
class ExcelRow:
    number: int
    link: str
    reviewed: str


@dataclass
class RawRow:
    number: int
    values: dict[str, str]
    links: dict[str, str]


@dataclass
class ParsedSheet:
    name: str
    headers: list[str]
    header_row: int
    rows: list[RawRow]
    suggested_link: str | None
    suggested_text: str | None


@dataclass
class SheetPreview:
    name: str
    headers: list[str]
    header_row: int
    suggested_link: str | None
    suggested_text: str | None
    data_rows: int
    sample: list[dict[str, str]]
    suggested: bool = False


@dataclass
class RowResult:
    row: int
    link: str
    reviewed: str
    extracted: str
    equality: float | None
    pages: int
    words_extracted: int
    words_reviewed: int
    chars_extracted: int
    chars_reviewed: int
    needs_ocr: bool
    error: str


def fold(text: str) -> str:
    decomposed = unicodedata.normalize("NFKD", text)
    plain = "".join(ch for ch in decomposed if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", " ", plain.lower()).strip()


def normalize_for_equality(text: str) -> str:
    """Un salto de linea cuenta como un espacio. No cambia letras ni puntuacion."""
    text = _ZERO_WIDTH.sub("", text or "")
    text = text.replace("\u00a0", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def equality_percent(reviewed: str, extracted: str) -> float:
    left = normalize_for_equality(reviewed)
    right = normalize_for_equality(extracted)
    if not left and not right:
        return 100.0
    if not left or not right:
        return 0.0
    ratio = difflib.SequenceMatcher(None, left, right, autojunk=False).ratio()
    return round(ratio * 100, 1)


def word_count(text: str) -> int:
    normalized = normalize_for_equality(text)
    if not normalized:
        return 0
    return len(normalized.split(" "))


def score_link(label: str) -> int:
    folded = fold(label)
    if not folded:
        return 0
    has_link = any(token in folded for token in ("liga", "url", "link", "pdf", "enlace", "hipervinculo"))
    if any(token in folded for token in ("transcripcion", "revisad", "contenido")) and not has_link:
        return 0
    for index, hint in enumerate(_LINK_HINTS):
        if hint in folded:
            return 100 - index
    return 0


def score_text(label: str) -> int:
    folded = fold(label)
    if not folded:
        return 0
    if any(token in folded for token in ("extraid", "ocr", "automatic")):
        return 0
    looks_like_link = any(token in folded for token in ("liga", "url", "link", "hipervinculo", "enlace"))
    if looks_like_link and "texto" not in folded and "transcripcion" not in folded:
        return 0
    for index, hint in enumerate(_TEXT_HINTS):
        if hint in folded:
            return 100 - index
    return 0


def suggest_columns(headers: list[str]) -> tuple[str | None, str | None]:
    link = max(headers, key=score_link, default=None)
    text = max(headers, key=score_text, default=None)
    if link is not None and score_link(link) <= 0:
        link = None
    if text is not None and score_text(text) <= 0:
        text = None
    if link and text and link == text:
        text = None
    return link, text


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def clean_link(value: str) -> str:
    value = value.strip().strip('"').strip("'")
    lowered = value.lower()
    if lowered.startswith(("http://", "https://")):
        return re.sub(r"\s+", "", value)
    return value


def looks_like_reference(value: str) -> bool:
    lowered = value.lower()
    if lowered.startswith(("http://", "https://", "file:")):
        return True
    if lowered.split("?", 1)[0].endswith(".pdf"):
        return True
    if re.match(r"^[a-zA-Z]:[\\/]", value) or value.startswith("\\\\"):
        return True
    return False


def _formula_hyperlink(value: object) -> str:
    if not isinstance(value, str) or not value.startswith("="):
        return ""
    match = _HYPERLINK_URL.search(value)
    if not match:
        return ""
    return clean_link(match.group("q") or match.group("s") or "")


def _formula_label(value: object) -> str:
    if not isinstance(value, str):
        return ""
    match = _HYPERLINK_LABEL.search(value)
    if not match:
        return ""
    return (match.group("q") or match.group("s") or "").strip()


def display_text(cell) -> str:
    value = cell.value
    if isinstance(value, str) and value.startswith("="):
        label = _formula_label(value)
        if label:
            return label
        hyperlink = getattr(cell, "hyperlink", None)
        if hyperlink is not None and getattr(hyperlink, "display", None):
            return str(hyperlink.display).strip()
        return ""
    return cell_text(value)


def link_from_cell(cell) -> str:
    target = ""
    hyperlink = getattr(cell, "hyperlink", None)
    if hyperlink is not None and hyperlink.target:
        target = clean_link(str(hyperlink.target))
    formula = _formula_hyperlink(cell.value)
    text = ""
    if not (isinstance(cell.value, str) and cell.value.startswith("=")):
        text = clean_link(cell_text(cell.value))
    for candidate in (target, formula, text):
        if candidate and looks_like_reference(candidate):
            return candidate
    return target or formula or text


def _unique_headers(raw: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    headers: list[str] = []
    for index, label in enumerate(raw, start=1):
        base = label.strip() or f"Columna {index}"
        count = seen.get(base, 0) + 1
        seen[base] = count
        headers.append(base if count == 1 else f"{base} ({count})")
    return headers


def _header_key(labels: list[str]) -> tuple[int, int, int]:
    # Una celda larga es transcripcion, no un encabezado.
    usable = [label if len(label) <= 80 else "" for label in labels]
    link = max((score_link(label) for label in usable), default=0)
    text = max((score_text(label) for label in usable), default=0)
    nonempty = sum(1 for label in usable if label.strip())
    return (1 if link and text else 0, link + text, nonempty)


def _row_labels(cells) -> list[str]:
    labels = [display_text(cell) for cell in cells]
    while labels and not labels[-1].strip():
        labels.pop()
    return labels


def read_workbook(path: Path) -> list[ParsedSheet]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    sheets: list[ParsedSheet] = []
    try:
        for worksheet in workbook.worksheets:
            scan_end = min(worksheet.max_row or 1, 10)
            scanned = list(worksheet.iter_rows(min_row=1, max_row=scan_end))
            best_index = 0
            best_key = (-1, -1, -1)
            best_labels: list[str] = []
            for offset, cells in enumerate(scanned):
                labels = _row_labels(cells)
                key = _header_key(labels)
                if labels and key > best_key:
                    best_key = key
                    best_index = offset
                    best_labels = labels
            if not best_labels:
                sheets.append(
                    ParsedSheet(
                        name=worksheet.title,
                        headers=[],
                        header_row=1,
                        rows=[],
                        suggested_link=None,
                        suggested_text=None,
                    )
                )
                continue

            headers = _unique_headers(best_labels)
            header_row = best_index + 1
            rows: list[RawRow] = []
            empty_run = 0
            last_row = min(worksheet.max_row or header_row, header_row + _MAX_DATA_ROWS)
            for cells in worksheet.iter_rows(min_row=header_row + 1, max_row=last_row):
                values: dict[str, str] = {}
                links: dict[str, str] = {}
                for header, cell in zip(headers, cells):
                    values[header] = display_text(cell)
                    links[header] = link_from_cell(cell)
                if not any(values.values()) and not any(links.values()):
                    empty_run += 1
                    if rows and empty_run >= 30:
                        break
                    continue
                empty_run = 0
                rows.append(RawRow(number=cells[0].row, values=values, links=links))
            link_column, text_column = suggest_columns(headers)
            sheets.append(
                ParsedSheet(
                    name=worksheet.title,
                    headers=headers,
                    header_row=header_row,
                    rows=rows,
                    suggested_link=link_column,
                    suggested_text=text_column,
                )
            )
    finally:
        workbook.close()
    if not sheets:
        raise ValueError("El Excel no tiene hojas.")
    return sheets


def preview_workbook(path: Path) -> list[SheetPreview]:
    sheets = read_workbook(path)
    best = max(
        range(len(sheets)),
        key=lambda index: (
            1 if sheets[index].suggested_link and sheets[index].suggested_text else 0,
            1 if sheets[index].rows else 0,
            len(sheets[index].rows),
        ),
    )
    previews: list[SheetPreview] = []
    for index, sheet in enumerate(sheets):
        sample = []
        for row in sheet.rows[:3]:
            item = {}
            for header in sheet.headers:
                text = row.values.get(header, "")
                link = row.links.get(header, "")
                shown = f"{text} → {link}" if link and link != text else (link or text)
                if len(shown) > 180:
                    shown = shown[:177] + "..."
                item[header] = shown
            sample.append(item)
        previews.append(
            SheetPreview(
                name=sheet.name,
                headers=sheet.headers,
                header_row=sheet.header_row,
                suggested_link=sheet.suggested_link,
                suggested_text=sheet.suggested_text,
                data_rows=len(sheet.rows),
                sample=sample,
                suggested=index == best,
            )
        )
    return previews


def load_rows(path: Path, sheet_name: str, link_column: str, text_column: str) -> list[ExcelRow]:
    sheets = read_workbook(path)
    sheet = next((item for item in sheets if item.name == sheet_name), None)
    if sheet is None:
        available = ", ".join(item.name for item in sheets) or "(ninguna)"
        raise ValueError(f"No existe la hoja {sheet_name!r}. Hojas: {available}")
    if link_column not in sheet.headers or text_column not in sheet.headers:
        available = ", ".join(sheet.headers) or "(sin encabezados)"
        raise ValueError(f"Elige columnas que existan en la hoja. Encabezados: {available}")
    if link_column == text_column:
        raise ValueError("La liga y la transcripcion tienen que ser columnas distintas.")
    rows: list[ExcelRow] = []
    for raw in sheet.rows:
        link = (raw.links.get(link_column) or raw.values.get(link_column) or "").strip()
        reviewed = raw.values.get(text_column, "")
        if not link and not reviewed.strip():
            continue
        rows.append(ExcelRow(number=raw.number, link=link, reviewed=reviewed))
    if len(rows) > _MAX_DATA_ROWS:
        raise ValueError(f"El Excel tiene demasiadas filas. El maximo es {_MAX_DATA_ROWS}.")
    return rows


def direct_download_url(url: str) -> str:
    drive = re.search(r"drive\.google\.com/file/d/([^/]+)", url)
    if drive:
        file_id = drive.group(1)
        return f"https://drive.google.com/uc?export=download&id={file_id}"
    if "dropbox.com" in url and "dl=0" in url:
        return url.replace("dl=0", "dl=1")
    return url


def resolve_local(link: str) -> Path | None:
    raw = link.strip().strip('"').strip("'")
    if raw.lower().startswith("file:"):
        parsed = urlparse(raw)
        raw = urllib.request.url2pathname(parsed.path)
        if os.name == "nt" and raw.startswith("\\") and len(raw) > 2 and raw[2] == ":":
            raw = raw.lstrip("\\")
    path = Path(raw).expanduser()
    if path.is_file():
        return path
    return None


def _reject_unless_pdf(path: Path) -> None:
    head = path.read_bytes()[:1024] if path.exists() else b""
    if b"%PDF-" not in head:
        path.unlink(missing_ok=True)
        raise ValueError("La liga no devolvio un PDF.")


def download_pdf(url: str, dest: Path, timeout: int = 90) -> None:
    request = urllib.request.Request(
        direct_download_url(url),
        headers={
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"
            ),
            "Accept": "application/pdf,*/*",
        },
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            dest.parent.mkdir(parents=True, exist_ok=True)
            total = 0
            with dest.open("wb") as stream:
                while True:
                    chunk = response.read(64 * 1024)
                    if not chunk:
                        break
                    total += len(chunk)
                    if total > _MAX_PDF_BYTES:
                        raise ValueError("El PDF excede 80 MB.")
                    stream.write(chunk)
        _reject_unless_pdf(dest)
    except urllib.error.HTTPError as exc:
        dest.unlink(missing_ok=True)
        raise ValueError(f"No se pudo descargar el PDF ({exc.code}).") from exc
    except urllib.error.URLError as exc:
        dest.unlink(missing_ok=True)
        raise ValueError(f"No se pudo abrir la liga: {exc.reason}") from exc
    except TimeoutError as exc:
        dest.unlink(missing_ok=True)
        raise ValueError("La descarga del PDF agoto el tiempo de espera.") from exc
    except ValueError:
        dest.unlink(missing_ok=True)
        raise


def materialize_pdf(link: str, dest: Path) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    cleaned = clean_link(link)
    if not cleaned:
        raise ValueError("La fila no tiene liga al PDF.")
    if cleaned.lower().startswith(("http://", "https://")):
        download_pdf(cleaned, dest)
        return
    local = resolve_local(cleaned)
    if local is None:
        raise ValueError(f"No se encontro el PDF: {cleaned}")
    shutil.copyfile(local, dest)
    _reject_unless_pdf(dest)


def process_rows(
    rows: list[ExcelRow],
    work_dir: Path,
    genre: str,
    header_percent: float,
    footer_percent: float,
    on_progress=None,
    on_result=None,
) -> list[RowResult]:
    header_percent = max(0.0, min(float(header_percent), 40.0))
    footer_percent = max(0.0, min(float(footer_percent), 40.0))
    pdf_dir = work_dir / "pdfs"
    text_dir = work_dir / "text"
    pdf_dir.mkdir(parents=True, exist_ok=True)
    text_dir.mkdir(parents=True, exist_ok=True)
    results: list[RowResult] = []
    total = len(rows)
    for index, row in enumerate(rows, start=1):
        if on_progress:
            on_progress(index, total, row, "descargando")
        try:
            if not row.link.strip():
                raise ValueError("La fila no tiene liga al PDF.")
            pdf_path = pdf_dir / f"fila_{row.number}.pdf"
            materialize_pdf(row.link, pdf_path)
            if on_progress:
                on_progress(index, total, row, "extrayendo")
            extracted_result = extract_one_pdf(
                pdf_path=pdf_path,
                output_dir=text_dir,
                base_input=pdf_dir,
                mode="human",
                write_pages_json=False,
                header_percent=header_percent,
                footer_percent=footer_percent,
                genre=genre,
            )
            extracted = Path(extracted_result.output_txt).read_text(encoding="utf-8")
            if on_progress:
                on_progress(index, total, row, "comparando")
            result = RowResult(
                row=row.number,
                link=row.link,
                reviewed=row.reviewed,
                extracted=extracted,
                equality=equality_percent(row.reviewed, extracted),
                pages=extracted_result.page_count,
                words_extracted=word_count(extracted),
                words_reviewed=word_count(row.reviewed),
                chars_extracted=len(extracted),
                chars_reviewed=len(row.reviewed),
                needs_ocr=extracted_result.needs_ocr,
                error="",
            )
        except Exception as exc:
            result = RowResult(
                row=row.number,
                link=row.link,
                reviewed=row.reviewed,
                extracted="",
                equality=None,
                pages=0,
                words_extracted=0,
                words_reviewed=word_count(row.reviewed),
                chars_extracted=0,
                chars_reviewed=len(row.reviewed),
                needs_ocr=False,
                error=str(exc),
            )
            if on_progress:
                on_progress(index, total, row, "error")
        results.append(result)
        if on_result:
            on_result(index, total, result)
    return results


def training_pairs(results: list[RowResult]) -> list[RowResult]:
    """Pares con texto extraido y transcripcion revisada, listos para entrenar."""
    pairs: list[RowResult] = []
    for result in results:
        if result.error or result.needs_ocr:
            continue
        if not normalize_for_equality(result.extracted) or not normalize_for_equality(result.reviewed):
            continue
        pairs.append(result)
    return pairs


def summarize(results: list[RowResult], genre: str) -> dict:
    scores = [item.equality for item in results if item.equality is not None]
    return {
        "rows": len(results),
        "compared": len(scores),
        "errors": sum(1 for item in results if item.error),
        "average": round(sum(scores) / len(scores), 1) if scores else None,
        "minimum": min(scores) if scores else None,
        "needs_ocr": sum(1 for item in results if item.needs_ocr),
        "training_pairs": len(training_pairs(results)),
        "genre": genre,
    }


_EXCEL_CELL_LIMIT = 32767


def _excel_text(value: str) -> str:
    if len(value) <= _EXCEL_CELL_LIMIT:
        return value
    return value[: _EXCEL_CELL_LIMIT - 1] + "…"


def _style_header(sheet, fill: PatternFill, font: Font) -> None:
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center")


def write_result_workbook(path: Path, results: list[RowResult], summary: dict) -> Path:
    workbook = Workbook()
    header_fill = PatternFill("solid", fgColor="0F766E")
    header_font = Font(color="FFFFFF", bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    training = workbook.active
    training.title = "Entrenamiento"
    training.append(["Fila", "Liga", "Texto extraido", "Transcripcion revisada", "Igualdad %"])
    _style_header(training, header_fill, header_font)
    for result in training_pairs(results):
        training.append(
            [
                result.row,
                _excel_text(result.link),
                _excel_text(result.extracted),
                _excel_text(result.reviewed),
                result.equality,
            ]
        )
        row_cells = training[training.max_row]
        row_cells[2].alignment = wrap
        row_cells[3].alignment = wrap
        if result.equality is not None:
            row_cells[4].number_format = "0.0"
    for index, width in enumerate([10, 42, 80, 80, 14], start=1):
        training.column_dimensions[get_column_letter(index)].width = width
    training.freeze_panes = "A2"
    training.auto_filter.ref = f"A1:E{max(training.max_row, 1)}"

    sheet = workbook.create_sheet("Resultados")
    headers = [
        "Fila",
        "Liga",
        "Igualdad %",
        "Palabras extraidas",
        "Palabras revisadas",
        "Paginas",
        "Requiere OCR",
        "Error",
        "Texto extraido",
        "Transcripcion revisada",
    ]
    sheet.append(headers)
    _style_header(sheet, header_fill, header_font)
    for result in results:
        sheet.append(
            [
                result.row,
                result.link,
                result.equality,
                result.words_extracted,
                result.words_reviewed,
                result.pages,
                "si" if result.needs_ocr else "",
                result.error,
                _excel_text(result.extracted),
                _excel_text(result.reviewed),
            ]
        )
        row_cells = sheet[sheet.max_row]
        for cell in row_cells[7:10]:
            cell.alignment = wrap
        if result.equality is not None:
            row_cells[2].number_format = "0.0"
    widths = [10, 42, 14, 20, 20, 12, 14, 36, 80, 80]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:J{max(sheet.max_row, 1)}"

    summary_sheet = workbook.create_sheet("Resumen")
    summary_sheet.append(["Concepto", "Valor"])
    labels = [
        ("Filas", summary.get("rows")),
        ("Comparadas", summary.get("compared")),
        ("Con error", summary.get("errors")),
        ("Igualdad promedio", summary.get("average")),
        ("Igualdad minima", summary.get("minimum")),
        ("Sin capa de texto", summary.get("needs_ocr")),
        ("Pares de entrenamiento", summary.get("training_pairs")),
        ("Genero", summary.get("genre")),
    ]
    for label, value in labels:
        summary_sheet.append([label, value])
    summary_sheet.column_dimensions["A"].width = 24
    summary_sheet.column_dimensions["B"].width = 24
    for cell in summary_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path
