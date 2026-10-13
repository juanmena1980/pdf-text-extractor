from __future__ import annotations

import json
import os
import shutil
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import fitz
from openpyxl import Workbook, load_workbook

from comparativo.compare import (
    RowResult,
    direct_download_url,
    equality_percent,
    load_rows,
    materialize_pdf,
    preview_workbook,
    process_rows,
    summarize,
    write_result_workbook,
)
from comparativo.server import make_server
from extract_pdf_text import extract_one_pdf


PHRASE = (
    "El alcalde confirmo que la obra del parque central terminara en marzo. "
    "Los vecinos pidieron que el texto de esta nota quede completo y legible."
)


def make_pdf(path: Path, text: str = PHRASE) -> None:
    document = fitz.open()
    page = document.new_page()
    page.insert_textbox(fitz.Rect(72, 280, 540, 520), text, fontsize=12)
    document.save(path)
    document.close()


def write_excel(path: Path, rows: list[tuple], headers=("Liga", "Transcripcion revisada")) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Notas"
    sheet.append(["Reporte del dia"])
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)
    workbook.close()


def post_multipart(url: str, fields: dict[str, str], filename: str, content: bytes) -> tuple[int, dict]:
    boundary = "----comparativoTestBoundary7f3a"
    chunks = []
    for key, value in fields.items():
        chunks.append(
            (
                f"--{boundary}\r\n"
                f'Content-Disposition: form-data; name="{key}"\r\n\r\n'
                f"{value}\r\n"
            ).encode("utf-8")
        )
    chunks.append(
        (
            f"--{boundary}\r\n"
            f'Content-Disposition: form-data; name="excel"; filename="{filename}"\r\n'
            f"Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n"
        ).encode("utf-8")
        + content
        + b"\r\n"
    )
    chunks.append(f"--{boundary}--\r\n".encode("utf-8"))
    body = b"".join(chunks)
    request = urllib.request.Request(
        url,
        data=body,
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
    )
    try:
        with urllib.request.urlopen(request) as response:
            return response.status, json.load(response)
    except urllib.error.HTTPError as exc:
        return exc.code, json.load(exc)


class EqualityTest(unittest.TestCase):
    def test_line_breaks_count_as_spaces(self) -> None:
        self.assertEqual(equality_percent("hola mundo", "hola\nmundo"), 100.0)
        self.assertEqual(equality_percent("hola mundo", "hola\r\n\r\nmundo"), 100.0)
        self.assertLess(equality_percent("hola mundo", "holamundo"), 100.0)

    def test_missing_text_lowers_the_percentage(self) -> None:
        reviewed = "El texto completo y legible de la nota."
        extracted = "El texto completo."
        self.assertLess(equality_percent(reviewed, extracted), 90.0)

    def test_accents_remain_part_of_the_comparison(self) -> None:
        self.assertLess(equality_percent("canción", "cancion"), 100.0)


class WorkbookTest(unittest.TestCase):
    def test_detects_headers_below_a_title_and_reads_hyperlinks(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "notas.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Notas"
            sheet.append(["Sintesis del dia"])
            sheet.append(["Medio", "Liga PDF", "Transcripción revisada"])
            sheet.append(["El Diario", "ver nota", "Texto revisado de la nota."])
            sheet.cell(3, 2).hyperlink = "https://ejemplo.test/nota.pdf"
            formula = sheet.cell(4, 2, '=HYPERLINK("https://ejemplo.test/otra.pdf","abrir")')
            sheet.cell(4, 1, "Radio")
            sheet.cell(4, 3, "Segunda transcripcion")
            empty = workbook.create_sheet("Vacia")
            empty["A1"] = "sin datos"
            del formula
            workbook.save(path)
            workbook.close()

            previews = preview_workbook(path)
            notes = next(item for item in previews if item.name == "Notas")
            self.assertTrue(notes.suggested)
            self.assertEqual(notes.header_row, 2)
            self.assertEqual(notes.suggested_link, "Liga PDF")
            self.assertEqual(notes.suggested_text, "Transcripción revisada")
            self.assertEqual(notes.data_rows, 2)

            rows = load_rows(path, "Notas", "Liga PDF", "Transcripción revisada")
            self.assertEqual(rows[0].link, "https://ejemplo.test/nota.pdf")
            self.assertEqual(rows[0].reviewed, "Texto revisado de la nota.")
            self.assertEqual(rows[1].link, "https://ejemplo.test/otra.pdf")
            self.assertEqual(rows[1].reviewed, "Segunda transcripcion")

    def test_drive_view_link_becomes_a_direct_download(self) -> None:
        url = direct_download_url("https://drive.google.com/file/d/abc123/view?usp=sharing")
        self.assertEqual(url, "https://drive.google.com/uc?export=download&id=abc123")


class HtmlLinkHandler(BaseHTTPRequestHandler):
    def do_GET(self) -> None:
        body = b"<html>esto no es un pdf</html>"
        self.send_response(200)
        self.send_header("Content-Type", "text/html")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, format: str, *args: object) -> None:
        return


class ProcessTest(unittest.TestCase):
    def test_html_response_is_rejected(self) -> None:
        server = ThreadingHTTPServer(("127.0.0.1", 0), HtmlLinkHandler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            port = server.server_address[1]
            with tempfile.TemporaryDirectory() as folder:
                dest = Path(folder) / "fila.pdf"
                with self.assertRaisesRegex(ValueError, "no devolvio un PDF"):
                    materialize_pdf(f"http://127.0.0.1:{port}/nota", dest)
                self.assertFalse(dest.exists())
        finally:
            server.shutdown()
            server.server_close()

    def test_compares_extracted_text_ignoring_line_breaks(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            pdf_path = root / "nota.pdf"
            make_pdf(pdf_path)
            extracted = extract_one_pdf(
                pdf_path=pdf_path,
                output_dir=root / "primera",
                base_input=root,
                mode="human",
                write_pages_json=False,
                header_percent=0,
                footer_percent=0,
                genre="generico",
            )
            actual = Path(extracted.output_txt).read_text(encoding="utf-8")
            self.assertIn("parque central", actual.lower())
            self.assertIn("completo y legible", actual.lower())
            reviewed = "\n".join(actual.split())
            excel_path = root / "notas.xlsx"
            write_excel(
                excel_path,
                [
                    (str(pdf_path), reviewed),
                    (str(pdf_path), "Esta transcripcion revisada no aparece en el pdf."),
                    (str(root / "no-existe.pdf"), "Texto que no se puede extraer."),
                ],
            )
            rows = load_rows(excel_path, "Notas", "Liga", "Transcripcion revisada")
            results = process_rows(
                rows,
                root / "trabajo",
                genre="generico",
                header_percent=0,
                footer_percent=0,
            )
            self.assertEqual(results[0].equality, 100.0)
            self.assertIn("parque central", results[0].extracted.lower())
            self.assertLess(results[1].equality, 80.0)
            self.assertTrue(results[2].error)
            self.assertIsNone(results[2].equality)


class ServerTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.temp = tempfile.TemporaryDirectory()
        os.environ["COMPARATIVO_DATA"] = cls.temp.name
        cls.server = make_server("127.0.0.1", 0)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        host, port = cls.server.server_address
        cls.base = f"http://{host}:{port}"

    @classmethod
    def tearDownClass(cls) -> None:
        cls.server.shutdown()
        cls.server.server_close()
        cls.temp.cleanup()

    def test_preview_job_and_download(self) -> None:
        root = Path(self.temp.name) / "caso"
        if root.exists():
            shutil.rmtree(root)
        root.mkdir()
        pdf_path = root / "nota.pdf"
        make_pdf(pdf_path)
        extracted = extract_one_pdf(
            pdf_path=pdf_path,
            output_dir=root / "base",
            base_input=root,
            mode="human",
            write_pages_json=False,
            header_percent=0,
            footer_percent=0,
            genre="nota_informativa",
        )
        actual = Path(extracted.output_txt).read_text(encoding="utf-8")
        self.assertIn("parque central", actual.lower())
        excel_path = root / "notas.xlsx"
        write_excel(excel_path, [(str(pdf_path), "\n".join(actual.split()))])

        status, preview = post_multipart(
            f"{self.base}/api/preview",
            {},
            "notas.xlsx",
            excel_path.read_bytes(),
        )
        self.assertEqual(status, 200)
        sheet = preview["sheets"][0]
        self.assertEqual(sheet["suggested_link"], "Liga")
        self.assertEqual(sheet["suggested_text"], "Transcripcion revisada")

        status, job = post_multipart(
            f"{self.base}/api/jobs",
            {
                "sheet": sheet["name"],
                "link_column": "Liga",
                "text_column": "Transcripcion revisada",
                "genre": "nota_informativa",
                "header_percent": "0",
                "footer_percent": "0",
            },
            "notas.xlsx",
            excel_path.read_bytes(),
        )
        self.assertEqual(status, 200, job)
        deadline = 0
        snapshot = {}
        while deadline < 40:
            with urllib.request.urlopen(f"{self.base}/api/jobs/{job['id']}") as response:
                snapshot = json.load(response)
            if snapshot["status"] != "running":
                break
            deadline += 1
            threading.Event().wait(0.25)
        self.assertEqual(snapshot["status"], "done", snapshot.get("error"))
        self.assertEqual(snapshot["results"][0]["equality"], 100.0)
        self.assertIn("completo y legible", snapshot["results"][0]["extracted"].lower())

        with urllib.request.urlopen(snapshot["download"] and f"{self.base}{snapshot['download']}") as response:
            payload = response.read()
        result_path = root / "resultado.xlsx"
        result_path.write_bytes(payload)
        workbook = load_workbook(result_path, data_only=True)
        self.assertEqual(workbook.sheetnames[0], "Entrenamiento")
        training = workbook["Entrenamiento"]
        self.assertEqual(
            [cell.value for cell in training[1]],
            ["Fila", "Liga", "Texto extraido", "Transcripcion revisada", "Igualdad %"],
        )
        self.assertEqual(training["E2"].value, 100)
        self.assertIn("parque central", str(training["C2"].value).lower())
        reviewed = " ".join(str(training["D2"].value).split())
        self.assertIn("completo y legible", reviewed.lower())
        self.assertEqual(workbook["Resultados"]["C2"].value, 100)
        self.assertIn("parque central", str(workbook["Resultados"]["I2"].value).lower())
        workbook.close()


class TrainingWorkbookTest(unittest.TestCase):
    def test_training_sheet_keeps_complete_pairs_only(self) -> None:
        ready = RowResult(
            row=2,
            link="C:/notas/lista.pdf",
            reviewed="Texto revisado completo y legible.",
            extracted="Texto extraido completo y legible.",
            equality=92.5,
            pages=1,
            words_extracted=5,
            words_reviewed=5,
            chars_extracted=34,
            chars_reviewed=35,
            needs_ocr=False,
            error="",
        )
        failed = RowResult(
            row=3,
            link="C:/notas/ausente.pdf",
            reviewed="Esta fila no tiene PDF.",
            extracted="",
            equality=None,
            pages=0,
            words_extracted=0,
            words_reviewed=5,
            chars_extracted=0,
            chars_reviewed=24,
            needs_ocr=False,
            error="No se encontro el PDF.",
        )
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "entrenamiento.xlsx"
            write_result_workbook(path, [ready, failed], summarize([ready, failed], "nota_informativa"))
            workbook = load_workbook(path, data_only=True)
            training = workbook["Entrenamiento"]
            self.assertEqual(training.max_row, 2)
            self.assertEqual(training["A2"].value, 2)
            self.assertEqual(training["C2"].value, "Texto extraido completo y legible.")
            self.assertEqual(training["D2"].value, "Texto revisado completo y legible.")
            self.assertEqual(training["E2"].value, 92.5)
            self.assertEqual(workbook["Resultados"].max_row, 3)
            self.assertEqual(workbook["Resumen"]["A8"].value, "Pares de entrenamiento")
            self.assertEqual(workbook["Resumen"]["B8"].value, 1)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
